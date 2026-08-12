import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ─── Dados extraídos de categorias.js ────────────────────────────────────────

GRUPOS = {
    'RECEITAS':      'Receitas',
    'FIXAS':         'Despesas Fixas',
    'CONSUMO':       'Consumo Mensal',
    'DIVIDAS':       'Dívidas',
    'INVESTIMENTOS': 'Investimentos',
    'INTERNO':       'Fluxo Interno',
}

CATEGORIAS_PADRAO = [
    # RECEITAS
    {'id': 'cat_trabalho_clt',    'nome': 'Trabalho CLT',             'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': True},
    {'id': 'salario',             'nome': 'Salário',                  'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_trabalho_clt'},
    {'id': 'ferias',              'nome': 'Férias',                   'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_trabalho_clt'},
    {'id': 'decimo_terceiro',     'nome': '13º Salário',              'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_trabalho_clt'},
    {'id': 'bonus_plr',           'nome': 'Bônus/PLR',                'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_trabalho_clt'},

    {'id': 'cat_renda_variavel',  'nome': 'Renda Variável',           'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': True},
    {'id': 'dividendos',          'nome': 'Dividendos',               'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_renda_variavel'},
    {'id': 'rec_investimentos',   'nome': 'Rendimentos',              'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_renda_variavel'},
    {'id': 'juros_rec',           'nome': 'Juros Recebidos',          'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_renda_variavel'},
    {'id': 'renda_extra',         'nome': 'Renda Extra',              'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_renda_variavel'},
    {'id': 'aluguel_rec',         'nome': 'Aluguel Recebido',         'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_renda_variavel'},

    {'id': 'cat_entradas_nr',     'nome': 'Entradas Não Recorrentes', 'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': True},
    {'id': 'reembolso',           'nome': 'Reembolso',                'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_entradas_nr'},
    {'id': 'emprestimo_rec',      'nome': 'Empréstimos Recebidos',    'grupo': 'Receitas',       'tipo': 'receita',  'isCategoria': False, 'categoria': 'cat_entradas_nr'},

    # DESPESAS FIXAS
    {'id': 'cat_moradia',         'nome': 'Moradia',                  'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'aluguel_prest',       'nome': 'Aluguel/Prestação',        'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_moradia'},
    {'id': 'condominio',          'nome': 'Condomínio',               'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_moradia'},
    {'id': 'iptu',                'nome': 'IPTU',                     'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_moradia'},
    {'id': 'custos_imoveis',      'nome': 'Custos Imóveis',           'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_moradia'},

    {'id': 'cat_utilidades',      'nome': 'Utilidades',               'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'energia',             'nome': 'Energia',                  'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_utilidades'},
    {'id': 'agua',                'nome': 'Água',                     'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_utilidades'},
    {'id': 'gas',                 'nome': 'Gás',                      'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_utilidades'},
    {'id': 'internet',            'nome': 'Internet',                 'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_utilidades'},
    {'id': 'telefone',            'nome': 'Telefone',                 'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_utilidades'},
    {'id': 'celular',             'nome': 'Celular',                  'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_utilidades'},

    {'id': 'cat_seguros',         'nome': 'Seguros',                  'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'seguro_vida',         'nome': 'Seguro de Vida',           'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_seguros'},
    {'id': 'seguro_carro',        'nome': 'Seguro de Carro',          'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_seguros'},
    {'id': 'seguro_res',          'nome': 'Seguro Residencial',       'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_seguros'},

    {'id': 'cat_educ_fixa',       'nome': 'Educação Fixa',            'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'educacao',            'nome': 'Educação',                 'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_educ_fixa'},
    {'id': 'faculdade_mba',       'nome': 'Faculdade/MBA',            'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_educ_fixa'},
    {'id': 'educacao_filhos',     'nome': 'Educação Filhos',          'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_educ_fixa'},
    {'id': 'assinaturas',         'nome': 'Assinaturas',              'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_educ_fixa'},

    {'id': 'cat_funcionario',     'nome': 'Funcionário',              'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'funcionario',         'nome': 'Funcionário',              'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_funcionario'},
    {'id': 'salario_diaria',      'nome': 'Salário/Diária',           'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_funcionario'},
    {'id': 'encargos',            'nome': 'Encargos',                 'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_funcionario'},

    {'id': 'cat_obrigacoes',      'nome': 'Obrigações',               'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'ipva',                'nome': 'IPVA',                     'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_obrigacoes'},
    {'id': 'anuidade',            'nome': 'Anuidade',                 'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_obrigacoes'},
    {'id': 'tarifas',             'nome': 'Tarifas',                  'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_obrigacoes'},
    {'id': 'impostos',            'nome': 'Impostos',                 'grupo': 'Despesas Fixas', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_obrigacoes'},

    # CONSUMO MENSAL
    {'id': 'cat_alimentacao',     'nome': 'Alimentação',              'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'supermercado',        'nome': 'Supermercado',             'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_alimentacao'},
    {'id': 'alimentacao_fora',    'nome': 'Alimentação Fora',         'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_alimentacao'},
    {'id': 'padaria',             'nome': 'Padaria',                  'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_alimentacao'},

    {'id': 'cat_transporte',      'nome': 'Transporte',               'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'gasolina',            'nome': 'Gasolina',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},
    {'id': 'combustivel',         'nome': 'Combustível',              'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},
    {'id': 'uber',                'nome': 'Uber/Apps',                'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},
    {'id': 'carro',               'nome': 'Carro',                    'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},
    {'id': 'manutencao',          'nome': 'Manutenção',               'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},
    {'id': 'estacionamento',      'nome': 'Estacionamento',           'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},
    {'id': 'aluguel_carro',       'nome': 'Aluguel de Carro',         'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_transporte'},

    {'id': 'cat_saude',           'nome': 'Saúde',                    'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'saude',               'nome': 'Plano de Saúde',           'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_saude'},
    {'id': 'farmacia',            'nome': 'Farmácia',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_saude'},
    {'id': 'drogaria',            'nome': 'Drogaria',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_saude'},
    {'id': 'consulta',            'nome': 'Consulta',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_saude'},

    {'id': 'cat_bem_estar',       'nome': 'Bem-estar',                'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'academia',            'nome': 'Academia',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_bem_estar'},
    {'id': 'personal',            'nome': 'Personal',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_bem_estar'},
    {'id': 'esporte',             'nome': 'Esporte',                  'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_bem_estar'},
    {'id': 'beleza',              'nome': 'Beleza',                   'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_bem_estar'},
    {'id': 'cosmeticos',          'nome': 'Cosméticos',               'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_bem_estar'},
    {'id': 'cuidados_pessoais',   'nome': 'Cuidados Pessoais',        'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_bem_estar'},

    {'id': 'cat_vestuario',       'nome': 'Vestuário',                'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'vestuario',           'nome': 'Vestuário',                'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_vestuario'},

    {'id': 'cat_lazer',           'nome': 'Lazer',                    'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'lazer',               'nome': 'Lazer',                    'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_lazer'},
    {'id': 'viagem',              'nome': 'Viagem',                   'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_lazer'},
    {'id': 'clube',               'nome': 'Clube',                    'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_lazer'},
    {'id': 'comemoracoes',        'nome': 'Comemorações',             'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_lazer'},
    {'id': 'festa_aniversario',   'nome': 'Festa de Aniversário',     'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_lazer'},

    {'id': 'cat_casa',            'nome': 'Casa',                     'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'casa',                'nome': 'Casa',                     'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_casa'},
    {'id': 'decoracao',           'nome': 'Decoração',                'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_casa'},
    {'id': 'reforma',             'nome': 'Reforma',                  'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_casa'},

    {'id': 'cat_compras',         'nome': 'Compras',                  'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'compras_online',      'nome': 'Compras Online',           'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_compras'},
    {'id': 'presente',            'nome': 'Presente',                 'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_compras'},
    {'id': 'papelaria',           'nome': 'Papelaria',                'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_compras'},

    {'id': 'cat_educ_var',        'nome': 'Educação Variável',        'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'livro_curso',         'nome': 'Livro/Curso',              'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_educ_var'},
    {'id': 'material_escolar',    'nome': 'Material Escolar',         'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_educ_var'},

    {'id': 'cat_pets',            'nome': 'Pets',                     'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'pets',                'nome': 'Pets',                     'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_pets'},

    {'id': 'cat_outros',          'nome': 'Outros',                   'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': True},
    {'id': 'doacoes',             'nome': 'Doações',                  'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_outros'},
    {'id': 'saque',               'nome': 'Saque',                    'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_outros'},
    {'id': 'outros',              'nome': 'Outros',                   'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_outros'},
    {'id': 'nao_lembro',          'nome': 'Não Lembro',               'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_outros'},
    {'id': 'sem_categoria',       'nome': 'Sem Categoria',            'grupo': 'Consumo Mensal', 'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_outros'},

    # DÍVIDAS
    {'id': 'cat_dividas',         'nome': 'Dívidas',                  'grupo': 'Dívidas',        'tipo': 'despesa',  'isCategoria': True},
    {'id': 'fin_cartao',          'nome': 'Financiamento de Cartão',  'grupo': 'Dívidas',        'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_dividas'},
    {'id': 'emprestimo',          'nome': 'Empréstimos',              'grupo': 'Dívidas',        'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_dividas'},
    {'id': 'juros_div',           'nome': 'Juros',                    'grupo': 'Dívidas',        'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_dividas'},

    # INVESTIMENTOS
    {'id': 'cat_investimentos',   'nome': 'Investimentos',            'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': True},
    {'id': 'investimento',        'nome': 'Aplicações',               'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'previdencia',         'nome': 'Previdência Privada',      'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'reserva',             'nome': 'Reserva de Emergência',    'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'custodia',            'nome': 'Custódia',                 'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'custos_op',           'nome': 'Custos Operacionais',      'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'iof',                 'nome': 'IOF',                      'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'ir',                  'nome': 'IR',                       'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},
    {'id': 'perdas',              'nome': 'Perdas',                   'grupo': 'Investimentos',  'tipo': 'despesa',  'isCategoria': False, 'categoria': 'cat_investimentos'},

    # FLUXO INTERNO
    {'id': 'cat_fluxo_interno',   'nome': 'Fluxo Interno',            'grupo': 'Fluxo Interno',  'tipo': 'neutro',   'isCategoria': True},
    {'id': 'entre_contas',        'nome': 'Entre Contas',             'grupo': 'Fluxo Interno',  'tipo': 'neutro',   'isCategoria': False, 'categoria': 'cat_fluxo_interno'},
]

# ─── Mapa id → nome (para resolver categoria pai) ────────────────────────────
id_to_nome = {c['id']: c['nome'] for c in CATEGORIAS_PADRAO}

# ─── Tipo mapeado ─────────────────────────────────────────────────────────────
TIPO_LABEL = {
    'receita': 'Receita',
    'despesa': 'Despesa',
    'neutro':  'Neutro',
}

# ─── Cores por Macro ──────────────────────────────────────────────────────────
GRUPO_CORES = {
    'Receitas':        {'bg': '1A5C2A', 'fg': 'FFFFFF'},   # verde escuro
    'Despesas Fixas':  {'bg': '8B1A1A', 'fg': 'FFFFFF'},   # vermelho escuro
    'Consumo Mensal':  {'bg': 'C04000', 'fg': 'FFFFFF'},   # laranja escuro
    'Dívidas':         {'bg': '6B0080', 'fg': 'FFFFFF'},   # roxo
    'Investimentos':   {'bg': '0D3D6B', 'fg': 'FFFFFF'},   # azul escuro
    'Fluxo Interno':   {'bg': '3D3D3D', 'fg': 'FFFFFF'},   # cinza escuro
}

TIPO_CORES = {
    'Receita': {'bg': 'D4EDDA', 'fg': '145221'},
    'Despesa': {'bg': 'F8D7DA', 'fg': '721C24'},
    'Neutro':  {'bg': 'E2E3E5', 'fg': '383D41'},
}

# ─── Filtrar apenas subcategorias ────────────────────────────────────────────
subcats = [c for c in CATEGORIAS_PADRAO if not c['isCategoria']]

# ─── Criar workbook ──────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Categorias B2IF'

# ─── Estilos comuns ───────────────────────────────────────────────────────────
thin = Side(style='thin', color='CCCCCC')
border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

medium = Side(style='medium', color='999999')
border_medium = Border(left=medium, right=medium, top=medium, bottom=medium)

def make_fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

# ─── Título principal ─────────────────────────────────────────────────────────
ws.merge_cells('A1:F1')
title_cell = ws['A1']
title_cell.value = 'B2IF — Categorias e Subcategorias Padrão (v3.0)'
title_cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
title_cell.fill = make_fill('1A2D3D')
title_cell.alignment = Alignment(horizontal='center', vertical='center')
title_cell.border = border_medium
ws.row_dimensions[1].height = 30

# ─── Subtítulo com contagem ───────────────────────────────────────────────────
ws.merge_cells('A2:F2')
sub_cell = ws['A2']
n_cats  = len([c for c in CATEGORIAS_PADRAO if c['isCategoria']])
n_subs  = len(subcats)
sub_cell.value = f'{n_subs} subcategorias  •  {n_cats} categorias intermediárias  •  6 grupos macro'
sub_cell.font = Font(name='Calibri', size=10, italic=True, color='555555')
sub_cell.fill = make_fill('F0F4F8')
sub_cell.alignment = Alignment(horizontal='center', vertical='center')
sub_cell.border = border_thin
ws.row_dimensions[2].height = 18

# ─── Cabeçalho ────────────────────────────────────────────────────────────────
headers = ['Macro (Grupo)', 'Categoria', 'Subcategoria', 'ID', 'Tipo', 'Nível']
col_widths = [22, 26, 28, 28, 12, 10]

for col_idx, (h, w) in enumerate(zip(headers, col_widths), start=1):
    cell = ws.cell(row=3, column=col_idx, value=h)
    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    cell.fill = make_fill('243F5C')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_medium
    ws.column_dimensions[get_column_letter(col_idx)].width = w

ws.row_dimensions[3].height = 24

# ─── Dados ────────────────────────────────────────────────────────────────────
current_row = 4
prev_grupo = None
prev_cat_id = None

for item in subcats:
    grupo     = item['grupo']
    cat_id    = item.get('categoria', '')
    cat_nome  = id_to_nome.get(cat_id, cat_id)
    sub_nome  = item['nome']
    sub_id    = item['id']
    tipo      = TIPO_LABEL.get(item['tipo'], item['tipo'])

    gcores = GRUPO_CORES.get(grupo, {'bg': 'FFFFFF', 'fg': '000000'})
    tcores = TIPO_CORES.get(tipo,  {'bg': 'FFFFFF', 'fg': '000000'})

    # ── Coluna A: Macro ───────────────────────────────────────────────────────
    cell_a = ws.cell(row=current_row, column=1, value=grupo)
    cell_a.font = Font(name='Calibri', size=10, bold=True, color=gcores['fg'])
    cell_a.fill = make_fill(gcores['bg'])
    cell_a.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell_a.border = border_thin

    # ── Coluna B: Categoria ───────────────────────────────────────────────────
    cell_b = ws.cell(row=current_row, column=2, value=cat_nome)
    cell_b.font = Font(name='Calibri', size=10, bold=True, color='2C3E50')
    # Alterna levemente o fundo por categoria
    cat_fill = 'EAF0F6' if prev_cat_id != cat_id else 'FFFFFF'
    cell_b.fill = make_fill(cat_fill)
    cell_b.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    cell_b.border = border_thin

    # ── Coluna C: Subcategoria ────────────────────────────────────────────────
    cell_c = ws.cell(row=current_row, column=3, value=sub_nome)
    cell_c.font = Font(name='Calibri', size=10, color='2C3E50')
    cell_c.fill = make_fill('FFFFFF')
    cell_c.alignment = Alignment(horizontal='left', vertical='center', indent=2)
    cell_c.border = border_thin

    # ── Coluna D: ID ──────────────────────────────────────────────────────────
    cell_d = ws.cell(row=current_row, column=4, value=sub_id)
    cell_d.font = Font(name='Courier New', size=9, color='555555')
    cell_d.fill = make_fill('F8F9FA')
    cell_d.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    cell_d.border = border_thin

    # ── Coluna E: Tipo ────────────────────────────────────────────────────────
    cell_e = ws.cell(row=current_row, column=5, value=tipo)
    cell_e.font = Font(name='Calibri', size=10, bold=True, color=tcores['fg'])
    cell_e.fill = make_fill(tcores['bg'])
    cell_e.alignment = Alignment(horizontal='center', vertical='center')
    cell_e.border = border_thin

    # ── Coluna F: Nível (sempre Subcategoria) ─────────────────────────────────
    cell_f = ws.cell(row=current_row, column=6, value='Subcategoria')
    cell_f.font = Font(name='Calibri', size=9, italic=True, color='888888')
    cell_f.fill = make_fill('F8F9FA')
    cell_f.alignment = Alignment(horizontal='center', vertical='center')
    cell_f.border = border_thin

    ws.row_dimensions[current_row].height = 18

    prev_grupo  = grupo
    prev_cat_id = cat_id
    current_row += 1

# ─── Segunda aba: Categorias intermediárias ───────────────────────────────────
ws2 = wb.create_sheet('Categorias (Nível 2)')

ws2.merge_cells('A1:E1')
t2 = ws2['A1']
t2.value = 'B2IF — Categorias Intermediárias (Nível 2)'
t2.font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
t2.fill = make_fill('243F5C')
t2.alignment = Alignment(horizontal='center', vertical='center')
t2.border = border_medium
ws2.row_dimensions[1].height = 28

headers2 = ['Macro (Grupo)', 'Categoria', 'ID', 'Tipo', 'Nível']
widths2  = [22, 30, 28, 12, 14]

for ci, (h, w) in enumerate(zip(headers2, widths2), start=1):
    cell = ws2.cell(row=2, column=ci, value=h)
    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    cell.fill = make_fill('1A5280')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_medium
    ws2.column_dimensions[get_column_letter(ci)].width = w

ws2.row_dimensions[2].height = 22

cats_nivel2 = [c for c in CATEGORIAS_PADRAO if c['isCategoria']]
for ri, cat in enumerate(cats_nivel2, start=3):
    grupo    = cat['grupo']
    cat_nome = cat['nome']
    cat_id   = cat['id']
    tipo     = TIPO_LABEL.get(cat['tipo'], cat['tipo'])
    gcores   = GRUPO_CORES.get(grupo, {'bg': 'FFFFFF', 'fg': '000000'})
    tcores   = TIPO_CORES.get(tipo,  {'bg': 'FFFFFF', 'fg': '000000'})

    ws2.cell(ri, 1, grupo).font     = Font(name='Calibri', size=10, bold=True, color=gcores['fg'])
    ws2.cell(ri, 1).fill            = make_fill(gcores['bg'])
    ws2.cell(ri, 1).alignment       = Alignment(horizontal='center', vertical='center')
    ws2.cell(ri, 1).border          = border_thin

    ws2.cell(ri, 2, cat_nome).font  = Font(name='Calibri', size=10, bold=True, color='2C3E50')
    ws2.cell(ri, 2).fill            = make_fill('EBF5FB')
    ws2.cell(ri, 2).alignment       = Alignment(horizontal='left', vertical='center', indent=1)
    ws2.cell(ri, 2).border          = border_thin

    ws2.cell(ri, 3, cat_id).font    = Font(name='Courier New', size=9, color='555555')
    ws2.cell(ri, 3).fill            = make_fill('F8F9FA')
    ws2.cell(ri, 3).alignment       = Alignment(horizontal='left', vertical='center', indent=1)
    ws2.cell(ri, 3).border          = border_thin

    ws2.cell(ri, 4, tipo).font      = Font(name='Calibri', size=10, bold=True, color=tcores['fg'])
    ws2.cell(ri, 4).fill            = make_fill(tcores['bg'])
    ws2.cell(ri, 4).alignment       = Alignment(horizontal='center', vertical='center')
    ws2.cell(ri, 4).border          = border_thin

    ws2.cell(ri, 5, 'Categoria').font      = Font(name='Calibri', size=9, italic=True, color='888888')
    ws2.cell(ri, 5).fill                   = make_fill('F8F9FA')
    ws2.cell(ri, 5).alignment              = Alignment(horizontal='center', vertical='center')
    ws2.cell(ri, 5).border                 = border_thin

    ws2.row_dimensions[ri].height = 18

# ─── Terceira aba: Legenda dos Macros ─────────────────────────────────────────
ws3 = wb.create_sheet('Legenda')

ws3.merge_cells('A1:C1')
t3 = ws3['A1']
t3.value = 'Legenda — Grupos Macro B2IF'
t3.font = Font(name='Calibri', size=13, bold=True, color='FFFFFF')
t3.fill = make_fill('243F5C')
t3.alignment = Alignment(horizontal='center', vertical='center')
t3.border = border_medium
ws3.row_dimensions[1].height = 28

for ci, h in enumerate(['Macro (Grupo)', 'Tipo predominante', 'Descrição'], start=1):
    cell = ws3.cell(row=2, column=ci, value=h)
    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    cell.fill = make_fill('1A5280')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_medium

ws3.column_dimensions['A'].width = 22
ws3.column_dimensions['B'].width = 18
ws3.column_dimensions['C'].width = 50
ws3.row_dimensions[2].height = 22

legenda = [
    ('Receitas',       'Receita', 'Entradas de dinheiro: salário CLT, renda variável, entradas não recorrentes'),
    ('Despesas Fixas', 'Despesa', 'Gastos mensais fixos: moradia, utilidades, seguros, educação fixa, obrigações'),
    ('Consumo Mensal', 'Despesa', 'Gastos variáveis do dia a dia: alimentação, transporte, saúde, lazer, compras'),
    ('Dívidas',        'Despesa', 'Pagamento de dívidas: financiamentos de cartão, empréstimos, juros'),
    ('Investimentos',  'Despesa', 'Aplicações e aportes: aplicações, previdência, reserva de emergência, custódia'),
    ('Fluxo Interno',  'Neutro',  'Transferências entre contas próprias — não impacta o resultado líquido'),
]

for ri, (grupo, tipo, descr) in enumerate(legenda, start=3):
    gcores = GRUPO_CORES.get(grupo, {'bg': 'FFFFFF', 'fg': '000000'})
    tcores = TIPO_CORES.get(tipo,  {'bg': 'FFFFFF', 'fg': '000000'})

    ws3.cell(ri, 1, grupo).font      = Font(name='Calibri', size=10, bold=True, color=gcores['fg'])
    ws3.cell(ri, 1).fill             = make_fill(gcores['bg'])
    ws3.cell(ri, 1).alignment        = Alignment(horizontal='center', vertical='center')
    ws3.cell(ri, 1).border           = border_thin

    ws3.cell(ri, 2, tipo).font       = Font(name='Calibri', size=10, bold=True, color=tcores['fg'])
    ws3.cell(ri, 2).fill             = make_fill(tcores['bg'])
    ws3.cell(ri, 2).alignment        = Alignment(horizontal='center', vertical='center')
    ws3.cell(ri, 2).border           = border_thin

    ws3.cell(ri, 3, descr).font      = Font(name='Calibri', size=10, color='2C3E50')
    ws3.cell(ri, 3).fill             = make_fill('F8F9FA')
    ws3.cell(ri, 3).alignment        = Alignment(horizontal='left', vertical='center', indent=1, wrap_text=True)
    ws3.cell(ri, 3).border           = border_thin
    ws3.row_dimensions[ri].height    = 20

# ─── Freeze panes e filtros ───────────────────────────────────────────────────
ws.freeze_panes  = 'A4'
ws2.freeze_panes = 'A3'
ws.auto_filter.ref  = f'A3:F{current_row - 1}'
ws2.auto_filter.ref = f'A2:E{2 + len(cats_nivel2)}'

# ─── Salvar ───────────────────────────────────────────────────────────────────
output = '/home/user/webapp/B2IF_Categorias_v3.xlsx'
wb.save(output)
print(f'Arquivo gerado: {output}')
print(f'  Subcategorias: {n_subs}')
print(f'  Categorias (nível 2): {n_cats}')
