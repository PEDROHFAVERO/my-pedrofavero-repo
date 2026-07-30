# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Configurar margens ────────────────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Inches(8.27)   # A4
section.page_height = Inches(11.69)
section.left_margin   = Inches(1.0)
section.right_margin  = Inches(1.0)
section.top_margin    = Inches(1.0)
section.bottom_margin = Inches(1.0)

# ── Paleta de cores ───────────────────────────────────────────────────────────
COR_TITULO    = RGBColor(0x1A, 0x56, 0xDB)   # azul
COR_EPICO     = RGBColor(0x0F, 0x76, 0x6E)   # verde-azul
COR_HISTORIA  = RGBColor(0x37, 0x41, 0x51)   # cinza escuro
COR_P0        = RGBColor(0xDC, 0x26, 0x26)   # vermelho
COR_P1        = RGBColor(0xEA, 0x58, 0x0C)   # laranja
COR_P2        = RGBColor(0xCA, 0x8A, 0x04)   # amarelo escuro
COR_P3        = RGBColor(0x16, 0xA3, 0x4A)   # verde
COR_OK        = RGBColor(0x16, 0xA3, 0x4A)   # verde
COR_TEXTO     = RGBColor(0x11, 0x18, 0x27)   # quase preto
COR_MUTED     = RGBColor(0x6B, 0x72, 0x80)   # cinza

FONT = "Calibri"

def set_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def add_heading(text, level=1, color=COR_TITULO):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    sizes = {1: 18, 2: 14, 3: 12}
    set_font(run, size=sizes.get(level, 11), bold=True, color=color)
    return p

def add_para(text="", size=10.5, bold=False, italic=False, color=None, indent=0, space_before=2, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent   = Inches(indent * 0.25)
    p.paragraph_format.space_before  = Pt(space_before)
    p.paragraph_format.space_after   = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic, color=color or COR_TEXTO)
    return p

def add_bullet(text, indent=1, color=None, size=10.5, bold=False):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Inches(indent * 0.3)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color or COR_TEXTO)
    return p

def add_separator():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run("─" * 72)
    set_font(run, size=8, color=COR_MUTED)

def prioridade_cor(p):
    return {
        "P0": COR_P0, "P1": COR_P1, "P2": COR_P2, "P3": COR_P3
    }.get(p, COR_TEXTO)

# ══════════════════════════════════════════════════════════════════════════════
# CAPA
# ══════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(40)
r = p.add_run("B2IF — BACKLOG DE PRODUTO")
set_font(r, size=24, bold=True, color=COR_TITULO)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Planejador Financeiro PF  |  Versão 1.0  |  Abril 2026")
set_font(r, size=12, italic=True, color=COR_MUTED)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# CONTEXTO
# ══════════════════════════════════════════════════════════════════════════════
add_heading("Contexto do Produto", 1)

dados = [
    ("Empresa",           "B2 Inteligência Financeira"),
    ("Planejadores",      "4"),
    ("Clientes hoje",     "~100 ativos"),
    ("Meta 6 meses",      "~200 clientes"),
    ("Meta 1 ano",        "~500 clientes"),
    ("Modelo",            "Ferramenta interna da B2, sem comercialização externa"),
    ("Plataforma atual",  "Desktop (análise completa)"),
    ("Plataforma futura", "Mobile (acompanhamento de gastos no dia a dia)"),
]
for chave, valor in dados:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    r1 = p.add_run(f"{chave}: ")
    set_font(r1, size=10.5, bold=True, color=COR_TEXTO)
    r2 = p.add_run(valor)
    set_font(r2, size=10.5, color=COR_TEXTO)

# ══════════════════════════════════════════════════════════════════════════════
# VISÃO
# ══════════════════════════════════════════════════════════════════════════════
add_heading("Visão de Produto", 1)
add_para(
    '"Ser a central financeira do cliente B2: onde ele entende onde está, vê para onde vai, '
    'acompanha seu progresso em tempo real e é motivado a evoluir — com o planejador como guia."',
    size=11, italic=True, color=COR_EPICO, indent=1
)

# ══════════════════════════════════════════════════════════════════════════════
# LEGENDA
# ══════════════════════════════════════════════════════════════════════════════
add_heading("Legenda de Prioridades", 2, COR_HISTORIA)
for p_label, desc, cor in [
    ("P0 — Crítico",  "Bloqueia operação ou crescimento. Executar imediatamente.", COR_P0),
    ("P1 — Alta",     "Impacto direto na experiência. Próximo ciclo.",             COR_P1),
    ("P2 — Média",    "Melhoria relevante. Pode aguardar janela adequada.",         COR_P2),
    ("P3 — Futuro",   "Visão de longo prazo. Planejar sem data firme.",             COR_P3),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    r1 = p.add_run(f"  {p_label}  ")
    set_font(r1, size=10.5, bold=True, color=cor)
    r2 = p.add_run(desc)
    set_font(r2, size=10.5, color=COR_TEXTO)

# ══════════════════════════════════════════════════════════════════════════════
# MVP CONCLUÍDO
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("MVP Atual — Concluído", 1, COR_OK)

mvp_items = [
    "Autenticação: Manager / Planejador / Cliente",
    "Painel Manager — gestão de planejadores",
    "Hub do Planejador — carteira de clientes",
    "Dashboard Financeiro — visão anual e mensal por macro",
    "Fluxo Financeiro — importação CSV / XLSM / PDF",
    "Categorização manual e automática por regras",
    "Planejador Financeiro — orçamento mensal com projeções",
    "Acesso de leitura para o cliente (login próprio)",
    "Gestão de acessos de cliente (criar / resetar / excluir)",
    "Modo cliente — visualização restrita sem edição",
    "Filtros encadeados (cascata) nas transações",
    "Deploy manual no Netlify",
]
for item in mvp_items:
    add_bullet(f"[OK]  {item}", color=COR_OK)

# ══════════════════════════════════════════════════════════════════════════════
# HELPER: bloco de história
# ══════════════════════════════════════════════════════════════════════════════
def add_historia(codigo, titulo, prioridade, narrativa, criterios, tarefas=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(f"  {codigo}  |  ")
    set_font(r1, size=11, bold=True, color=COR_HISTORIA)
    r2 = p.add_run(titulo)
    set_font(r2, size=11, bold=True, color=COR_HISTORIA)
    r3 = p.add_run(f"  [{prioridade}]")
    set_font(r3, size=11, bold=True, color=prioridade_cor(prioridade))

    add_para(narrativa, size=10, italic=True, indent=1, color=COR_MUTED)

    add_para("Critérios de aceite:", size=10, bold=True, indent=1)
    for c in criterios:
        add_bullet(f"[ ]  {c}", indent=2, size=10)

    if tarefas:
        add_para("Tarefas técnicas:", size=10, bold=True, indent=1)
        for t in tarefas:
            add_bullet(t, indent=2, size=10, color=COR_MUTED)

# ══════════════════════════════════════════════════════════════════════════════
# EPICO 1
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Épico 1 — Ingestão e Tratamento de Dados", 1, COR_EPICO)
add_para("Eliminar o gargalo manual de preparar extratos via IA externa antes de subir no sistema.",
         size=10.5, italic=True, color=COR_MUTED)

add_historia(
    "E1-H1", "Upload inteligente de extratos brutos", "P0",
    "Como planejador, quero enviar o extrato bruto do cliente (PDF, foto, CSV de qualquer banco) "
    "direto no sistema, para que ele trate, normalize e importe automaticamente sem precisar de IA externa.",
    [
        "Aceitar PDF de qualquer banco (Itaú, Bradesco, Nubank, BB, XP, Inter...)",
        "Aceitar foto/imagem de extrato via OCR",
        "Aceitar CSV em diferentes formatos bancários",
        "Detectar data, valor, descrição e tipo (crédito/débito) automaticamente",
        "Exibir pré-visualização antes de confirmar importação",
        "Permitir correção manual antes de salvar",
        "Relatório de linhas não reconhecidas",
    ],
    [
        "Integrar OCR (Google Vision ou Tesseract) para imagens",
        "Parser universal de PDF bancário com fallback para OCR",
        "Pipeline de normalização de colunas por heurística",
        "Tela de pré-visualização com edição inline",
        "Endpoint de upload na Edge Function",
    ]
)

add_historia(
    "E1-H2", "Detecção automática de banco/formato", "P1",
    "Como planejador, quero que o sistema detecte automaticamente de qual banco é o extrato, "
    "para não precisar escolher o formato manualmente.",
    [
        "Identificar banco pelo cabeçalho/layout do arquivo",
        "Aplicar parser específico por banco automaticamente",
        "Fallback para parser genérico quando banco não reconhecido",
        "Lembrar qual banco cada conta do cliente usa",
    ]
)

add_historia(
    "E1-H3", "Integração Open Finance", "P3",
    "Como planejador, quero conectar a conta bancária do cliente via Open Finance, "
    "para que as transações sejam importadas automaticamente sem upload manual.",
    [
        "Conectar via API Pluggy ou Belvo",
        "Sincronização periódica automática (diária ou em tempo real)",
        "Planejador aprova/revisa antes de categorizar",
        "Histórico de sincronizações por conta",
    ]
)
add_para("Nota: Dependente de regulação Bacen e consentimento do cliente (LGPD).",
         size=9.5, italic=True, indent=1, color=COR_MUTED)

# ══════════════════════════════════════════════════════════════════════════════
# EPICO 2
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Épico 2 — Acompanhamento de Gastos em Tempo Real", 1, COR_EPICO)
add_para("Capturar gastos do cliente no dia a dia e atualizar o sistema automaticamente, "
         "dando visibilidade ao cliente de quanto já gastou das suas metas mensais.",
         size=10.5, italic=True, color=COR_MUTED)

add_historia(
    "E2-H1", "Captura de gasto via WhatsApp", "P0",
    "Como cliente, quero enviar foto de cupom fiscal, áudio ou texto no WhatsApp, "
    "para que o gasto seja registrado automaticamente no meu perfil financeiro.",
    [
        "Bot WhatsApp recebe mensagem do cliente (Twilio, Z-API ou Evolution API)",
        "Processar foto de cupom: extrair estabelecimento, valor, data",
        'Processar texto livre: "gastei 40 reais no lanche" → transação criada',
        "Processar áudio: transcrever e interpretar",
        "IA categoriza automaticamente com base nas regras do cliente",
        'Confirmação de volta: "R$ 40,00 em Alimentação Fora registrado. Restam R$ 960,00 na meta de Lanche."',
        "Planejador pode revisar/corrigir gastos enviados pelo cliente",
    ],
    [
        "Configurar webhook WhatsApp Business API",
        "Integrar OCR para cupons fiscais",
        "Integrar Whisper (OpenAI) ou Google Speech para áudio",
        "Integrar GPT para interpretação de texto livre",
        "Endpoint de entrada de transação avulsa por cliente_id",
        "Resposta automática com saldo restante da categoria",
    ]
)

add_historia(
    "E2-H2", 'Painel mobile "Hoje" (PWA)', "P1",
    "Como cliente, quero ver no celular quanto já gastei hoje e no mês por categoria, "
    "para tomar decisões conscientes no momento do gasto.",
    [
        "Tela mobile com barras de progresso por categoria",
        "Destaque para categorias próximas ou acima da meta",
        "Últimas transações do dia",
        "Saldo disponível por macro (Consumo, Fixas, etc.)",
        "Funciona como PWA (sem precisar de app store)",
    ]
)

add_historia(
    "E2-H3", "Notificações proativas de meta", "P2",
    "Como cliente, quero receber alertas quando estou chegando perto do limite de uma categoria, "
    "para ajustar meu comportamento antes de estourar.",
    [
        "Alerta ao atingir 80% da meta de uma categoria",
        "Alerta ao ultrapassar a meta",
        "Canal configurável: WhatsApp, push (PWA) ou e-mail",
        "Planejador define quais categorias têm alerta por cliente",
    ]
)

# ══════════════════════════════════════════════════════════════════════════════
# EPICO 3
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Épico 3 — Gamificação e Progresso do Cliente", 1, COR_EPICO)
add_para("Dar ao cliente uma visão lúdica e motivadora do seu progresso financeiro, "
         "transformando o planejamento em uma jornada com marcos visíveis.",
         size=10.5, italic=True, color=COR_MUTED)

add_historia(
    "E3-H1", "Objetivos e Metas Financeiras", "P1",
    "Como planejador, quero cadastrar objetivos financeiros para o cliente "
    "(reserva de emergência, viagem, imóvel, aposentadoria), "
    "para que ele veja para onde está indo e o que já conquistou.",
    [
        "Criar objetivo: nome, valor alvo, prazo, ícone/foto",
        "Vincular objetivo a uma categoria de Investimentos",
        "Calcular progresso automaticamente com base nos aportes realizados",
        "Exibir % concluído e valor atual vs. alvo",
        "Calcular aporte mensal necessário para atingir no prazo",
        "Alertar quando aporte do mês não foi feito",
    ]
)

add_historia(
    "E3-H2", "Dashboard Gamificado", "P2",
    "Como cliente, quero ver meu progresso financeiro de forma visual e motivadora, "
    "para me sentir engajado com minha evolução.",
    [
        "Linha do tempo de evolução patrimonial desde o início do planejamento",
        'Conquistas desbloqueadas: "Primeiro mês sem estourar meta", "3 meses com saldo positivo"...',
        "Score financeiro mensal 0–100 (adimplência às metas, saldo positivo, aportes realizados)",
        'Comparativo mês a mês: "Em Janeiro você gastou R$ 500 a mais em Consumo que em Dezembro"',
        "Animações simples ao atingir uma meta",
    ]
)

add_historia(
    "E3-H3", "Relatório Mensal Automático pós-sessão", "P1",
    "Como planejador, quero gerar e enviar automaticamente um relatório pós-sessão, "
    "para que o cliente saia da reunião com resumo claro do que foi discutido e os próximos passos.",
    [
        "PDF com: panorama do mês anterior, desvios da meta, evolução do score, objetivos em andamento",
        "Campo para o planejador adicionar: ajustes propostos, deveres de casa, alvos do próximo mês",
        'Envio automático por e-mail ou WhatsApp ao clicar "Fechar sessão"',
        "Cliente acessa histórico de relatórios dentro do sistema",
        "Layout profissional com branding B2",
    ]
)

add_historia(
    "E3-H4", "Simulador de Aposentadoria e Objetivos de Longo Prazo", "P2",
    "Como planejador, quero realizar simulações de aposentadoria e objetivos de longo prazo "
    "diretamente no sistema, para que o cliente visualize cenários e tome decisões baseadas em dados.",
    [
        "Simulador de aposentadoria: idade atual, idade desejada, patrimônio atual, aporte mensal, taxa → patrimônio projetado",
        "Modo usufruto total (gastar tudo) e usufruto parcial (preservar capital)",
        "Projeção de renda mensal na aposentadoria",
        'Simulador de objetivo: "Quero comprar um imóvel de R$ 500k em 5 anos. Quanto preciso poupar por mês?"',
        "Gráfico de projeção com e sem inflação",
        "Salvar simulações no perfil do cliente",
    ]
)

# ══════════════════════════════════════════════════════════════════════════════
# EPICO 4
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Épico 4 — IA Integrada ao Sistema", 1, COR_EPICO)
add_para("Um assistente de IA que conversa com planejadores e clientes, "
         "interpreta os dados do sistema e oferece insights automáticos.",
         size=10.5, italic=True, color=COR_MUTED)

add_historia(
    "E4-H1", "Assistente IA para o Planejador", "P2",
    "Como planejador, quero perguntar à IA sobre os dados de um cliente, "
    "para ter insights rápidos sem precisar analisar tabelas manualmente.",
    [
        "Chat embutido no sistema (por cliente ou visão geral)",
        "Contexto da IA inclui: transações, categorias, planejamento, objetivos e histórico",
        "Respostas com dados reais do sistema (não genéricas)",
        "Sugestões proativas de anomalias detectadas",
        "IA pode gerar rascunho de relatório pós-sessão",
    ]
)

add_historia(
    "E4-H2", "Assistente IA para o Cliente", "P2",
    "Como cliente, quero perguntar ao assistente sobre minha situação financeira, "
    "para entender meus dados sem precisar esperar a próxima sessão com o planejador.",
    [
        "Chat disponível na área do cliente (desktop e mobile)",
        "Responde apenas com base nos dados do próprio cliente",
        "Tom motivador e didático (não técnico)",
        "Não sugere produtos financeiros (compliance)",
    ]
)

add_historia(
    "E4-H3", "Transcrição e Resumo Automático de Sessões", "P3",
    "Como planejador, quero gravar e transcrever automaticamente a sessão com o cliente via Google Meet, "
    "para ter o resumo salvo no perfil do cliente sem precisar tomar notas manualmente.",
    [
        "Integração com Google Meet API para captura de transcrição",
        "IA gera resumo estruturado: tópicos discutidos, decisões tomadas, próximos passos",
        "Resumo salvo no perfil do cliente com data da sessão",
        "Planejador pode editar antes de salvar",
        "Histórico de sessões acessível no sistema",
    ]
)

# ══════════════════════════════════════════════════════════════════════════════
# EPICO 5
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Épico 5 — Integração Google Workspace", 1, COR_EPICO)
add_para("Central de produtividade do planejador: agenda, sessões, documentos e comunicação dentro do sistema.",
         size=10.5, italic=True, color=COR_MUTED)

add_historia(
    "E5-H1", "Agenda Google integrada ao sistema", "P2",
    "Como planejador, quero ver e gerenciar minha agenda do Google Calendar dentro do sistema, "
    "para não precisar alternar entre ferramentas durante o atendimento.",
    [
        "Visualização semanal/mensal da agenda no painel do planejador",
        "Criar, editar e cancelar compromissos",
        "Vincular evento a um cliente do sistema",
        "Notificação 15 min antes da sessão",
        "Link automático do Meet inserido no evento",
    ]
)

add_historia(
    "E5-H2", "Google Drive — documentos do cliente", "P3",
    "Como planejador, quero acessar e salvar documentos do cliente diretamente no Google Drive pelo sistema, "
    "para que tudo fique centralizado no perfil do cliente.",
    [
        "Pasta no Drive criada automaticamente ao criar novo cliente",
        "Upload de documentos (contratos, declarações, extratos) pelo sistema",
        "Visualização de arquivos do Drive dentro do perfil do cliente",
        "Relatórios pós-sessão salvos automaticamente na pasta do cliente",
    ]
)

add_historia(
    "E5-H3", "Gmail — comunicação centralizada", "P3",
    "Como planejador, quero enviar e-mails para o cliente diretamente do sistema, "
    "para que toda comunicação fique registrada no histórico do cliente.",
    [
        "Enviar e-mail usando conta Gmail do planejador via OAuth",
        "Templates pré-definidos: pós-sessão, lembrete de sessão, alerta de meta",
        "Histórico de e-mails enviados no perfil do cliente",
        "Resposta do cliente notifica o planejador no sistema",
    ]
)

# ══════════════════════════════════════════════════════════════════════════════
# EPICO 6
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Épico 6 — Infraestrutura e Escalabilidade", 1, COR_EPICO)
add_para("Garantir que o sistema suporte o crescimento para 500 clientes com estabilidade e segurança.",
         size=10.5, italic=True, color=COR_MUTED)

add_historia(
    "E6-H1", "Migração para ambiente de produção", "P0",
    "Como time B2, quero ter um ambiente de produção separado e estável, "
    "para que os dados reais dos clientes fiquem isolados do ambiente de desenvolvimento.",
    [
        "Supabase produção criado e configurado",
        "Schema, RLS e Edge Functions migrados",
        "Variáveis de ambiente configuradas no Netlify",
        "Domínio próprio configurado",
        "Backup automático habilitado no Supabase",
    ]
)

add_historia(
    "E6-H2", "Deploy automatizado via GitHub", "P2",
    "Como time B2, quero que o deploy seja feito automaticamente ao fazer push na branch main, "
    "para que novas versões cheguem em produção sem processo manual.",
    [
        "Repositório GitHub conectado ao Netlify",
        "Branch main → deploy automático em produção",
        "Branch develop → deploy automático em staging",
        "Notificação de sucesso/falha no deploy",
    ]
)

add_historia(
    "E6-H3", "LGPD — Consentimento e gestão de dados", "P1",
    "Como empresa B2, quero ter controles mínimos de LGPD implementados, "
    "para que estejamos em conformidade ao crescer a base de clientes.",
    [
        "Tela de aceite de termos de uso e política de privacidade no primeiro login do cliente",
        "Possibilidade de exportar todos os dados de um cliente (direito de portabilidade)",
        "Possibilidade de excluir completamente um cliente e todos seus dados",
        "Log de acesso: quem acessou qual perfil e quando",
    ]
)

add_historia(
    "E6-H4", "Monitoramento e alertas de erro", "P2",
    "Como time B2, quero ser alertado quando algo quebra em produção, "
    "para corrigir antes que o planejador ou cliente perceba.",
    [
        "Integrar Sentry (ou similar) para captura de erros de frontend",
        "Alertas por e-mail quando erro crítico ocorrer",
        "Log de erros nas Edge Functions do Supabase",
    ]
)

# ══════════════════════════════════════════════════════════════════════════════
# ROADMAP
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Roadmap por Marcos", 1)

marcos = [
    ("Marco 0", "MVP Desktop", "CONCLUÍDO", "Março 2026",
     ["Sistema no Netlify com acesso de cliente", "4 planejadores, ~100 clientes"],
     COR_OK),
    ("Marco 1", "Produção Estável", "Abr/Mai 2026", "4 planejadores operando com segurança em produção",
     ["Migração Supabase produção  (E6-H1)  P0",
      "Relatório pós-sessão PDF + envio  (E3-H3)  P1",
      "LGPD básico  (E6-H3)  P1",
      "Deploy automatizado GitHub  (E6-H2)  P2"],
     COR_P0),
    ("Marco 2", "Ingestão Inteligente de Dados", "Mai/Jun 2026",
     "Reduzir tempo de onboarding de extrato de ~30min para ~5min",
     ["Upload extrato bruto de qualquer banco  (E1-H1)  P0",
      "Detecção automática de banco/formato  (E1-H2)  P1",
      "Pré-visualização e correção antes de importar  (E1-H1)  P0"],
     COR_P0),
    ("Marco 3", "Gastos em Tempo Real", "Jul/Ago 2026",
     "~200 clientes | Cliente acompanha gastos fora da sessão mensal",
     ['Bot WhatsApp (foto, texto, áudio)  (E2-H1)  P0',
      'PWA mobile painel "Hoje"  (E2-H2)  P1',
      "Notificações de meta (80% e 100%)  (E2-H3)  P2"],
     COR_P1),
    ("Marco 4", "Gamificação e Objetivos", "Set/Out 2026",
     "~300 clientes | Engajamento e retenção do cliente",
     ["Objetivos e metas financeiras  (E3-H1)  P1",
      "Dashboard gamificado: score, conquistas  (E3-H2)  P2",
      "Simulador de aposentadoria  (E3-H4)  P2"],
     COR_P1),
    ("Marco 5", "IA e Google Workspace", "Q1 2027",
     "~500 clientes | Eficiência operacional máxima",
     ["Assistente IA para o planejador  (E4-H1)  P2",
      "Assistente IA para o cliente  (E4-H2)  P2",
      "Agenda Google integrada  (E5-H1)  P2",
      "Transcrição automática de sessões  (E4-H3)  P3",
      "Google Drive por cliente  (E5-H2)  P3"],
     COR_P2),
    ("Marco 6", "Open Finance", "Mid 2027",
     "Importação automática bancária sem upload manual",
     ["Integração Open Finance via Pluggy/Belvo  (E1-H3)  P3",
      "Sincronização automática por conta  (E1-H3)  P3"],
     COR_P3),
]

for codigo, titulo, periodo, meta, itens, cor in marcos:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(f"{codigo}  —  ")
    set_font(r1, size=12, bold=True, color=COR_HISTORIA)
    r2 = p.add_run(titulo)
    set_font(r2, size=12, bold=True, color=cor)
    r3 = p.add_run(f"   [{periodo}]")
    set_font(r3, size=10, italic=True, color=COR_MUTED)

    add_para(f"Meta: {meta}", size=10, italic=True, indent=1, color=COR_MUTED)
    for item in itens:
        add_bullet(item, indent=2, size=10)

# ══════════════════════════════════════════════════════════════════════════════
# BACKLOG PRIORIZADO
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Backlog Priorizado — Ordem de Execução", 1)

rows = [
    (1,  "Migração Supabase produção",           "E6-H1", "M1", "P0"),
    (2,  "Upload inteligente de extratos",        "E1-H1", "M2", "P0"),
    (3,  "Bot WhatsApp captura de gastos",        "E2-H1", "M3", "P0"),
    (4,  "Relatório pós-sessão PDF",              "E3-H3", "M1", "P1"),
    (5,  "LGPD básico",                           "E6-H3", "M1", "P1"),
    (6,  "Detecção automática de banco",          "E1-H2", "M2", "P1"),
    (7,  'PWA mobile painel "Hoje"',              "E2-H2", "M3", "P1"),
    (8,  "Objetivos e metas financeiras",         "E3-H1", "M4", "P1"),
    (9,  "Deploy automatizado GitHub",            "E6-H2", "M1", "P2"),
    (10, "Notificações de meta",                  "E2-H3", "M3", "P2"),
    (11, "Dashboard gamificado",                  "E3-H2", "M4", "P2"),
    (12, "Simulador de aposentadoria",            "E3-H4", "M4", "P2"),
    (13, "Monitoramento e alertas de erro",       "E6-H4", "—",  "P2"),
    (14, "Assistente IA para o planejador",       "E4-H1", "M5", "P2"),
    (15, "Assistente IA para o cliente",          "E4-H2", "M5", "P2"),
    (16, "Agenda Google integrada",               "E5-H1", "M5", "P2"),
    (17, "Transcrição automática de sessões",     "E4-H3", "M5", "P3"),
    (18, "Google Drive por cliente",              "E5-H2", "M5", "P3"),
    (19, "Gmail centralizado",                    "E5-H3", "M5", "P3"),
    (20, "Open Finance",                          "E1-H3", "M6", "P3"),
]

table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'
hdr = table.rows[0].cells
for i, h in enumerate(["#", "Item", "Épico", "Marco", "Prioridade"]):
    hdr[i].text = h
    run = hdr[i].paragraphs[0].runs[0]
    set_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    tc = hdr[i]._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '1A56DB')
    tcPr.append(shd)

for num, item, epico, marco, prio in rows:
    row_cells = table.add_row().cells
    row_cells[0].text = str(num)
    row_cells[1].text = item
    row_cells[2].text = epico
    row_cells[3].text = marco
    row_cells[4].text = prio
    cor_p = prioridade_cor(prio)
    for ci, cell in enumerate(row_cells):
        for para in cell.paragraphs:
            for run in para.runs:
                set_font(run, size=10, bold=(ci == 4), color=cor_p if ci == 4 else COR_TEXTO)

# ══════════════════════════════════════════════════════════════════════════════
# PERGUNTAS EM ABERTO
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_heading("Perguntas em Aberto", 1)
add_para("A definir antes de cada marco:", size=10.5, bold=True)

abertas = [
    ("Marco 3", "Qual provedor de WhatsApp Business? (Twilio, Z-API, Evolution API)"),
    ("Marco 3", "PWA ou app nativo (React Native)?"),
    ("Marco 4", "Fórmula do score financeiro — definir com equipe B2"),
    ("Marco 4", "Lista de conquistas/badges — definir com equipe B2"),
    ("Marco 5", "Qual modelo de IA? (OpenAI GPT-4, Gemini, Claude)"),
    ("Marco 5", "Conta Google corporativa da B2 para integração OAuth"),
    ("Marco 6", "Fluxo de consentimento Open Finance com o cliente"),
]
for marco, pergunta in abertas:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.3)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(f"[{marco}]  ")
    set_font(r1, size=10.5, bold=True, color=COR_TITULO)
    r2 = p.add_run(pergunta)
    set_font(r2, size=10.5, color=COR_TEXTO)

# ══════════════════════════════════════════════════════════════════════════════
# RODAPÉ
# ══════════════════════════════════════════════════════════════════════════════
add_separator()
add_para("Documento vivo — atualizar a cada marco concluído.", size=9, italic=True, color=COR_MUTED)
add_para("Próxima revisão sugerida: após conclusão do Marco 1 (Produção Estável).", size=9, italic=True, color=COR_MUTED)

# ── Salvar ────────────────────────────────────────────────────────────────────
doc.save("/home/user/webapp/b2if-pf/BACKLOG_B2IF.docx")
print("ok")
